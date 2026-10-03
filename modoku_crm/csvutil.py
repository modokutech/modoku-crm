"""Shared CSV-export helper — every module's list page can export its
current (filtered) rows as a CSV file in a couple of lines, reusing the same
RFC 6266/5987-safe filename header the PDF downloads already use.
"""
import csv
import io
import re

from flask import Response

from .docutil import content_disposition


def csv_response(filename, header, rows):
    """filename should already end in .csv. rows: an iterable of iterables
    (each item becomes one row's cells, in header order)."""
    buf = io.StringIO()
    writer = csv.writer(buf)
    writer.writerow(header)
    writer.writerows(rows)
    # A UTF-8 BOM so Excel (which otherwise guesses ANSI/Windows-1252) opens
    # accented or non-ASCII text correctly instead of mangling it.
    body = "﻿" + buf.getvalue()
    return Response(
        body, mimetype="text/csv; charset=utf-8",
        headers={"Content-Disposition": content_disposition(filename)},
    )


def class_export_name(session_row, suffix):
    """Fix122: Microsoft_Excel_Basic_2026-10-05_participants - the first
    three words of the course and the class start date, so exports from
    different classes don't all land as the same name in Downloads. No
    extension; table_response adds .csv or .xlsx."""
    words = re.findall(r"[A-Za-z0-9]+", session_row["course_title"] or "")[:3]
    parts = ["_".join(words), session_row["start_date"] or "", suffix]
    return "_".join(p for p in parts if p)


def table_response(name, header, rows, fmt="csv"):
    """Fix123: one export, two formats. fmt "xlsx" gives an Excel workbook
    (bold frozen header, sized columns, text kept as text so IC numbers
    keep their leading zeros); anything else gives the usual CSV. Floats
    are money here, so they show with 2 decimals in both."""
    if fmt != "xlsx":
        rows = ([f"{c:.2f}" if isinstance(c, float) else c for c in row] for row in rows)
        return csv_response(name + ".csv", header, rows)
    from openpyxl import Workbook
    from openpyxl.styles import Font
    from openpyxl.utils import get_column_letter

    wb = Workbook()
    ws = wb.active
    ws.title = "Export"
    ws.append(header)
    for cell in ws[1]:
        cell.font = Font(bold=True)
    ws.freeze_panes = "A2"
    widths = [len(str(h)) for h in header]
    for row in rows:
        ws.append(list(row))
        for i, c in enumerate(row):
            widths[i] = max(widths[i], len(str(c)) if c is not None else 0)
    for row in ws.iter_rows(min_row=2):
        for cell in row:
            if isinstance(cell.value, str):
                cell.number_format = "@"
            elif isinstance(cell.value, float):
                cell.number_format = "#,##0.00"
    for i, w in enumerate(widths, start=1):
        ws.column_dimensions[get_column_letter(i)].width = min(w + 2, 50)
    buf = io.BytesIO()
    wb.save(buf)
    return Response(
        buf.getvalue(),
        mimetype="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
        headers={"Content-Disposition": content_disposition(name + ".xlsx")},
    )
